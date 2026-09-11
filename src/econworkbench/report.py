# -*- coding: utf-8 -*-
"""
report.py — EconWorkbench report module: results CSV -> journal-ready three-line tables
Input : one or more crosscheck-contract CSVs (columns: term,estimate,se,pvalue)
Output: LaTeX (booktabs three-line), Word (docx three-line), PNG preview
"""
import sys, os, csv, argparse

NOTES = "Standard errors in parentheses.  *** p<0.01, ** p<0.05, * p<0.1"


def build_parser():
    ap = argparse.ArgumentParser(
        prog="econ-report",
        description="Turn result CSVs (term,estimate,se,pvalue) into journal-ready three-line tables.")
    ap.add_argument("csvs", nargs="+")
    ap.add_argument("--title", default="")
    ap.add_argument("--out", default="table")
    ap.add_argument("--format", default="tex,docx,png")
    ap.add_argument("--decimals", type=int, default=3)
    return ap


def load(path):
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            rows.append((r["term"].strip(), float(r["estimate"]), float(r["se"]), float(r["pvalue"])))
    return rows


def stars(p):
    return "***" if p < 0.01 else ("**" if p < 0.05 else ("*" if p < 0.1 else ""))


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    args = build_parser().parse_args(argv)

    models = []
    for p in args.csvs:
        models.append((os.path.splitext(os.path.basename(p))[0], load(p)))
    terms = [t for t, *_ in models[0][1]]
    d = args.decimals
    fmts = args.format.split(",")

    # ---------- LaTeX ----------
    if "tex" in fmts:
        tex = ["\\begin{table}[htbp]", "\\centering"]
        if args.title:
            tex.append("\\caption{" + args.title + "}")
        tex.append("\\begin{tabular}{l" + "c" * len(models) + "}")
        tex.append("\\toprule")
        tex.append(" & " + " & ".join(f"({i+1})" for i in range(len(models))) + " \\\\")
        tex.append("\\midrule")
        for t in terms:
            cells = [f"{row[1]:.{d}f}{stars(row[3])}"
                     if (row := next((r for r in rows if r[0] == t), None)) else ""
                     for _, rows in models]
            tex.append(f"{t} & " + " & ".join(cells) + " \\\\")
            ses = [f"({row[2]:.{d}f})" if (row := next((r for r in rows if r[0] == t), None)) else ""
                   for _, rows in models]
            if any(ses):
                tex.append(" & " + " & ".join(ses) + " \\\\")
        tex.append("\\midrule")
        tex.append(" & " + " & ".join(m[0] for m in models) + " \\\\")
        tex.append("\\bottomrule")
        tex.append("\\end{tabular}")
        tex.append("\\begin{flushleft}\\footnotesize " + NOTES.replace("<", "$<$") + "\\end{flushleft}")
        tex.append("\\end{table}")
        with open(args.out + ".tex", "w", encoding="utf-8") as f:
            f.write("\n".join(tex) + "\n")
        print("wrote", args.out + ".tex")

    # ---------- Word (three-line style) ----------
    if "docx" in fmts:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)

        if args.title:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(args.title)
            r.font.bold = True

        tbl = doc.add_table(rows=1 + 2 * len(terms) + 1, cols=1 + len(models))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        def set_borders(cell, top=None, bottom=None):
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge, sz in (("top", top), ("bottom", bottom)):
                if sz:
                    el = OxmlElement(f"w:{edge}")
                    el.set(qn("w:val"), "single")
                    el.set(qn("w:sz"), str(sz))
                    el.set(qn("w:color"), "000000")
                    borders.append(el)
            tcPr.append(borders)

        def put(cell, text, bold=False, center=True):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.bold = bold
            r.font.size = Pt(10.5)

        hdr = tbl.rows[0].cells
        for j in range(1 + len(models)):
            put(hdr[j], f"({j})" if j else "", bold=True)
            set_borders(hdr[j], top=12, bottom=6)

        for i, t in enumerate(terms):
            rc = tbl.rows[1 + 2 * i].cells
            put(rc[0], t, center=False)
            for j, (name, rows) in enumerate(models):
                row = next((r for r in rows if r[0] == t), None)
                put(rc[j + 1], (f"{row[1]:.{d}f}{stars(row[3])}") if row else "")
            sc = tbl.rows[2 + 2 * i].cells
            put(sc[0], "")
            for j, (name, rows) in enumerate(models):
                row = next((r for r in rows if r[0] == t), None)
                put(sc[j + 1], f"({row[2]:.{d}f})" if row else "")

        last = tbl.rows[-1].cells
        put(last[0], "", center=False)
        for j, (name, _) in enumerate(models):
            put(last[j + 1], name)
            set_borders(last[j + 1], bottom=12)
        set_borders(last[0], bottom=12)

        note_p = doc.add_paragraph()
        nr = note_p.add_run(NOTES)
        nr.font.size = Pt(8.5)

        doc.save(args.out + ".docx")
        print("wrote", args.out + ".docx")

    # ---------- PNG preview ----------
    if "png" in fmts:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        n_rows = 1 + 2 * len(terms) + 1
        cell_text = []
        for t in terms:
            coef = [f"{row[1]:.{d}f}{stars(row[3])}" if (row := next((r for r in rows if r[0] == t), None)) else ""
                    for _, rows in models]
            ses = [f"({row[2]:.{d}f})" if (row := next((r for r in rows if r[0] == t), None)) else ""
                   for _, rows in models]
            cell_text.append([t] + coef)
            cell_text.append([""] + ses)
        header = [""] + [f"({j+1})" for j in range(len(models))]
        footer = [""] + [name for name, _ in models]

        fig, ax = plt.subplots(figsize=(2.2 + 1.9 * len(models), 0.34 * n_rows + 1.1))
        ax.set_xlim(0, 1 + len(models))
        ax.set_ylim(-1.3, n_rows + 1)
        ax.axis("off")
        if args.title:
            ax.set_title(args.title, fontsize=12, pad=14)

        def draw_row(y, vals, bold=False, size=10):
            for j, v in enumerate(vals):
                ax.text(j + 0.02, y, v, fontsize=size,
                        fontweight="bold" if bold else "normal",
                        ha="left", va="center", family="serif")

        draw_row(n_rows + 0.35, header, bold=True)
        for i, row_vals in enumerate(cell_text):
            draw_row(n_rows - 0.55 - i, row_vals)
        draw_row(0.15, footer)
        ax.plot([0, 1 + len(models)], [n_rows + 0.75, n_rows + 0.75], color="k", lw=1.4)
        ax.plot([0, 1 + len(models)], [n_rows + 0.05, n_rows + 0.05], color="k", lw=0.7)
        ax.plot([0, 1 + len(models)], [-0.2, -0.2], color="k", lw=1.4)
        ax.text(0, -0.75, NOTES, fontsize=8, ha="left", va="center", family="serif")
        fig.savefig(args.out + ".png", dpi=200, bbox_inches="tight")
        print("wrote", args.out + ".png")

    return 0


if __name__ == "__main__":
    sys.exit(main())
